"""Your own resume, in any common format, structured by one Claude call.

Text extraction is local and per format; structuring is one call per ingest,
not per job. The output is the same `ResumeProfile` the reconciler has always
consumed, so nothing downstream knows which format the resume came in.

The lens summaries are NOT extracted: they come from `persona.yaml`, because
how you want to be positioned is yours to write, not a model's to infer.
"""

from __future__ import annotations

from pathlib import Path

from pydantic import Field

from .matching.scorer import DEFAULT_MODEL, Scorer
from .models import Project, ResumeJob, ResumeProfile, Skill, Strict, Summaries
from .parse_linkedin_pdf import extract_text as _pdf_text

SUPPORTED = (".pdf", ".docx", ".txt", ".md")


class ResumeError(RuntimeError):
    """The resume could not be read or structured. Fatal to ingest."""


class _Extraction(Strict):
    skills: list[Skill] = Field(default_factory=list)
    jobs: list[ResumeJob] = Field(default_factory=list)
    projects: list[Project] = Field(default_factory=list)


SYSTEM = """\
You convert a resume into structured data. Copy, never invent: every bullet,
skill and project must come from the resume text, in its own wording.

- jobs: one per position, most recent first. `org` is the employer exactly as
  written; `title` is the job title. `date_placeholder` is true only when the
  resume shows no dates for that position. Bullets are the accomplishment
  lines under it.
- Tag each bullet's `audience`: `design` if it evidences design, UX or
  research work; `eng` if it evidences engineering or building; `both` if it
  evidences both or neither clearly. Never use `untagged`.
- skills: each listed skill, tagged `design` or `eng`.
- projects: named side projects or published work, if any, with a one-line
  `detail`.
"""


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ResumeError(
            f"{path.name}: {suffix or 'no extension'} is not supported - "
            f"export your resume as PDF or DOCX ({', '.join(SUPPORTED)} work)")
    if not path.is_file():
        raise ResumeError(f"resume not found at {path}")
    if suffix == ".pdf":
        return _pdf_text(path)
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.extend(cell.text for cell in row.cells)
        return "\n".join(lines)
    return path.read_text(encoding="utf-8")


def parse_file(path: Path, summaries: Summaries, client=None,
               model: str = DEFAULT_MODEL) -> ResumeProfile:
    text = extract_text(Path(path)).strip()
    if not text:
        raise ResumeError(f"{Path(path).name}: no text could be extracted - "
                          "if it is a scanned image, export a text PDF or DOCX")
    if client is None:
        import anthropic
        client = anthropic.Anthropic()
    out = Scorer(client, model=model)._call(
        SYSTEM, f"Structure this resume:\n\n{text}", _Extraction, Path(path).name)
    return ResumeProfile(summaries=summaries, skills=out.skills, jobs=out.jobs,
                         projects=out.projects)
