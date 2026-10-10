"""The resume parser: per-format text extraction, then one structured call.

The client is a fake - no test makes a live API call.
"""
from pathlib import Path

import pytest

from jobhunt.models import ResumeBullet, ResumeJob, Skill, Summaries
from jobhunt.parse_resume import ResumeError, _Extraction, extract_text, parse_file
from jobhunt.persona import load_persona

ROOT = Path(__file__).resolve().parents[1]
RESUME = ("Sam Example\nProduct Designer\n\nSpokewell Repair - Co-Owner\n"
          "- Designed the booking flow end to end.\n")
SUMMARIES = Summaries(design="A designer.", eng="A builder.")


class FakeBlock:
    def __init__(self, parsed): self.type = "text"; self.parsed_output = parsed
class FakeMsg:
    def __init__(self, parsed):
        self.content = [FakeBlock(parsed)]; self.stop_reason = "end_turn"
        self.usage = None
class FakeMessages:
    def __init__(self, parsed): self._p = parsed; self.kwargs = None
    def parse(self, **kw):
        self.kwargs = kw
        return FakeMsg(self._p)
class FakeClient:
    def __init__(self, parsed): self.messages = FakeMessages(parsed)


EXTRACTION = _Extraction(
    skills=[Skill(text="Prototyping", audience="design")],
    jobs=[ResumeJob(org="Spokewell Repair", title="Co-Owner", date_placeholder=False,
                    bullets=[ResumeBullet(text="Designed the booking flow end to end.",
                                          audience="design")])])


def test_txt_and_md_are_read_as_text(tmp_path):
    for suffix in (".txt", ".md"):
        path = tmp_path / f"resume{suffix}"
        path.write_text(RESUME, encoding="utf-8")
        assert "Spokewell Repair" in extract_text(path)


def test_docx_paragraphs_and_tables_are_extracted(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("Spokewell Repair - Co-Owner")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Prototyping"
    path = tmp_path / "resume.docx"
    doc.save(str(path))
    text = extract_text(path)
    assert "Spokewell Repair" in text and "Prototyping" in text


@pytest.mark.parametrize("suffix", [".doc", ".pages", ".odt", ".rtf", ""])
def test_unsupported_formats_say_what_to_export_instead(tmp_path, suffix):
    path = tmp_path / f"resume{suffix}"
    path.write_bytes(b"x")
    with pytest.raises(ResumeError, match="export your resume as PDF or DOCX"):
        extract_text(path)


def test_a_missing_file_names_its_path(tmp_path):
    with pytest.raises(ResumeError, match="not found"):
        extract_text(tmp_path / "resume.docx")


def test_an_empty_resume_is_loud_not_an_empty_profile(tmp_path):
    path = tmp_path / "resume.txt"
    path.write_text("   \n", encoding="utf-8")
    with pytest.raises(ResumeError, match="no text"):
        parse_file(path, SUMMARIES, client=FakeClient(EXTRACTION))


def test_parse_file_sends_the_text_and_takes_summaries_from_the_persona(tmp_path):
    path = tmp_path / "resume.txt"
    path.write_text(RESUME, encoding="utf-8")
    client = FakeClient(EXTRACTION)
    profile = parse_file(path, SUMMARIES, client=client)
    assert "Designed the booking flow" in client.messages.kwargs["messages"][0]["content"]
    assert profile.summaries == SUMMARIES
    assert profile.jobs[0].org == "Spokewell Repair"
    assert profile.jobs[0].bullets[0].audience == "design"


def test_the_shipped_example_persona_loads():
    persona = load_persona(ROOT / "persona.example.yaml")
    assert persona.summaries.design and persona.summaries.eng and persona.positioning


def test_a_missing_persona_says_how_to_create_one(tmp_path):
    with pytest.raises(FileNotFoundError, match="persona.example.yaml"):
        load_persona(tmp_path / "persona.yaml")
